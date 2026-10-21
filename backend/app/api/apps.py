"""Ứng dụng AI: Hỏi tài liệu, Đọc hoá đơn → Excel, Soạn báo giá, Biên bản họp, Phân tích doanh thu.

Mọi lệnh gọi AI đi qua services/llm.py (app="apps") để tính hạn mức, chọn model và ghi nhật ký.
"""
from __future__ import annotations

import base64
import io
import json
import logging
import re
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..db import get_db
from ..deps import web_user
from ..models import Branch, Job
from ..services import llm, rag, storage
from ..services.jsonutil import parse_json_object, to_number
from ..services.principal import Principal
from ..services.router import _plain, classify
from .common import raise_http

router = APIRouter(prefix="/api/apps", tags=["apps"])
log = logging.getLogger("hub.apps")


def _read_upload(file: UploadFile, max_bytes: int, label: str) -> bytes:
    data = file.file.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise HTTPException(413, f"File quá lớn (tối đa {label}).")
    if not data:
        raise HTTPException(400, "File trống.")
    return data


def _call(db: Session, p: Principal, feature: str, messages: list[dict], **kw) -> llm.ChatResult:
    try:
        return llm.chat(db, p, app="apps", feature=feature, messages=messages, **kw)
    except llm.GatewayError as e:
        raise_http(e)
        raise  # để trình kiểm tra kiểu hiểu hàm luôn ném lỗi


def _search(db: Session, p: Principal, query: str, *, k: int, public_only: bool, feature: str):
    try:
        return rag.search(db, query, k=k, public_only=public_only, p=p, app="apps", feature=feature)
    except llm.GatewayError as e:
        raise_http(e)
        raise


# ================= Hỏi tài liệu =================
ASK_SYSTEM = (
    "Bạn là trợ lý nội bộ của Thi Lốp (chuỗi dịch vụ lốp, ắc quy, chăm sóc và sửa chữa ô tô). "
    "Chỉ trả lời dựa trên các đoạn trích tài liệu được cung cấp. Ghi nguồn bằng số trong ngoặc vuông, ví dụ [1] hoặc [2][3], "
    "ngay sau ý lấy từ đoạn đó. Nếu các đoạn trích không có thông tin cần thiết, nói rõ: "
    "\"Tài liệu hiện có không nêu thông tin này.\" và không tự suy đoán, không bịa giá hay chính sách. "
    "Trả lời bằng tiếng Việt, ngắn gọn, rõ ràng."
)


class AskIn(BaseModel):
    question: str


@router.post("/ask")
def ask(body: AskIn, p: Principal = Depends(web_user), db: Session = Depends(get_db)):
    q = body.question.strip()
    if not q:
        raise HTTPException(400, "Chưa nhập câu hỏi.")
    if len(q) > 4000:
        raise HTTPException(400, "Câu hỏi quá dài (tối đa 4.000 ký tự).")
    results = _search(db, p, q, k=6, public_only=False, feature="ask")
    if not results:
        return {"answer": "Chưa có tài liệu nào sẵn sàng để tra cứu. Nhờ quản lý tải tài liệu lên mục Tài liệu trước nhé.",
                "sources": [], "cost_vnd": 0}
    user = f"Các đoạn trích tài liệu:\n\n{rag.format_context(results)}\n\nCâu hỏi: {q}"
    r = _call(db, p, "ask", [{"role": "system", "content": ASK_SYSTEM}, {"role": "user", "content": user}],
              task_hint=classify([{"role": "user", "content": q}]))
    sources = [{"document_id": d.id, "title": d.title, "page": c.page,
                "excerpt": (c.text[:197] + "...") if len(c.text) > 200 else c.text} for c, d, _ in results]
    return {"answer": r.content, "sources": sources, "cost_vnd": float(r.cost_vnd)}


# ================= Đọc hoá đơn → Excel =================
IMAGE_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}
INVOICE_PROMPT = (
    "Đọc hoá đơn / phiếu nhập hàng này và trả về DUY NHẤT một JSON đúng mẫu, không thêm chữ nào khác:\n"
    '{"fields": {"supplier": "tên nhà cung cấp", "invoice_no": "số hoá đơn", "date": "dd/mm/yyyy", '
    '"total": số tổng thanh toán, "tax": số tiền thuế}, '
    '"items": [{"name": "tên hàng", "quantity": số, "unit_price": số, "amount": số}]}\n'
    "Quy tắc: số tiền ghi dạng số nguyên VND, không dấu chấm phân cách, không chữ 'đ'. "
    "Thông tin nào không đọc được thì để null. Không tự bịa thêm dòng hàng."
)


def parse_invoice(text: str) -> dict:
    """Chuẩn hoá JSON hoá đơn do AI trả về (chịu được khung ``` và số dạng chuỗi)."""
    data = parse_json_object(text) or {}
    f = data.get("fields") if isinstance(data.get("fields"), dict) else {}
    fields = {
        "supplier": (str(f.get("supplier")).strip() if f.get("supplier") else ""),
        "invoice_no": (str(f.get("invoice_no")).strip() if f.get("invoice_no") else ""),
        "date": (str(f.get("date")).strip() if f.get("date") else ""),
        "total": to_number(f.get("total")),
        "tax": to_number(f.get("tax")),
    }
    items = []
    for it in data.get("items") or []:
        if not isinstance(it, dict):
            continue
        name = str(it.get("name") or "").strip()
        qty, price, amount = to_number(it.get("quantity")), to_number(it.get("unit_price")), to_number(it.get("amount"))
        if amount is None and qty is not None and price is not None:
            amount = qty * price
        if name or amount is not None:
            items.append({"name": name, "quantity": qty, "unit_price": price, "amount": amount})
    return {"fields": fields, "items": items}


def invoice_xlsx(fields: dict, items: list[dict]) -> bytes:
    """Tạo file Excel: thông tin hoá đơn, bảng hàng (dòng tiêu đề in đậm), các dòng cộng."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoa don"
    bold = Font(bold=True)
    money = "#,##0"
    ws.append(["Nhà cung cấp", fields.get("supplier") or ""])
    ws.append(["Số hoá đơn", fields.get("invoice_no") or ""])
    ws.append(["Ngày", fields.get("date") or ""])
    for r in range(1, 4):
        ws.cell(r, 1).font = bold
    ws.append([])
    ws.append(["STT", "Tên hàng", "Số lượng", "Đơn giá", "Thành tiền"])
    for c in ws[ws.max_row]:
        c.font = bold
    for i, it in enumerate(items, 1):
        ws.append([i, it.get("name") or "", it.get("quantity"), it.get("unit_price"), it.get("amount")])
        ws.cell(ws.max_row, 4).number_format = money
        ws.cell(ws.max_row, 5).number_format = money
    subtotal = sum(it.get("amount") or 0 for it in items)
    tax = fields.get("tax")
    total = fields.get("total")
    if total is None:
        total = subtotal + (tax or 0)
    ws.append([])
    for label, val in (("Cộng tiền hàng", subtotal), ("Thuế", tax), ("Tổng thanh toán", total)):
        ws.append(["", label, "", "", val])
        ws.cell(ws.max_row, 2).font = bold
        ws.cell(ws.max_row, 5).font = bold
        ws.cell(ws.max_row, 5).number_format = money
    for col, width in zip("ABCDE", (14, 46, 12, 16, 18), strict=True):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.post("/invoice")
def invoice(file: UploadFile = File(...), p: Principal = Depends(web_user), db: Session = Depends(get_db)):
    filename = file.filename or "hoa-don"
    ext = Path(filename).suffix.lower()
    if ext not in IMAGE_TYPES and ext != ".pdf":
        raise HTTPException(400, "Chỉ nhận ảnh (JPG, PNG, WEBP) hoặc PDF.")
    data = _read_upload(file, 10 * 1024 * 1024, "10 MB")
    if ext == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((pg.extract_text() or "") for pg in reader.pages[:10]).strip()
        except Exception as e:  # noqa: BLE001
            raise HTTPException(400, f"Không mở được file PDF: {e}") from e
        if len(text) < 20:
            raise HTTPException(400, "PDF này là ảnh scan, không có chữ để đọc. Vui lòng chụp ảnh hoá đơn (JPG/PNG) rồi tải lên lại.")
        messages = [{"role": "user", "content": f"{INVOICE_PROMPT}\n\nNội dung hoá đơn (chữ lấy từ PDF):\n{text[:20000]}"}]
        r = _call(db, p, "invoice", messages, task_hint="easy", use_cache=False)
    else:
        url = f"data:{IMAGE_TYPES[ext]};base64,{base64.b64encode(data).decode()}"
        messages = [{"role": "user", "content": [{"type": "text", "text": INVOICE_PROMPT},
                                                  {"type": "image_url", "image_url": {"url": url}}]}]
        r = _call(db, p, "invoice", messages, task_hint="vision", use_cache=False)
    parsed = parse_invoice(r.content)
    rel = storage.save_bytes(invoice_xlsx(parsed["fields"], parsed["items"]), Path(filename).stem + ".xlsx", folder="exports")
    return {**parsed, "xlsx_url": f"/api/files/{rel}", "cost_vnd": float(r.cost_vnd)}


# ================= Soạn báo giá =================
QUOTE_SYSTEM = (
    "Bạn soạn tin nhắn báo giá gửi khách cho Thi Lốp — chuỗi dịch vụ lốp, ắc quy, chăm sóc và sửa chữa ô tô. "
    "Giọng lịch sự, ấm áp, xưng 'Thi Lốp' hoặc 'em', gọi khách là 'anh/chị' (hoặc tên khách nếu có). "
    "Bố cục: lời chào; tóm tắt xe và nhu cầu; từng hạng mục kèm giá; dịch vụ đi kèm / bảo hành nếu tài liệu có nêu; "
    "thông tin cơ sở (địa chỉ, điện thoại, giờ mở cửa) nếu được cung cấp; lời mời đặt lịch. "
    "TUYỆT ĐỐI không bịa giá, khuyến mãi hay cam kết: chỉ dùng giá có trong phần 'Thông tin tham khảo'; "
    "hạng mục nào chưa có giá thì ghi [GIÁ]. Dưới 200 chữ, dùng gạch đầu dòng cho các hạng mục, không dùng markdown đậm/nghiêng."
)


class QuoteIn(BaseModel):
    customer: str | None = None
    vehicle: str
    tire_size: str | None = None
    services: list[str] = []
    branch_id: int | None = None
    notes: str | None = None


def branch_info(b: Branch | None) -> str:
    if b is None:
        return ""
    parts = [f"Cơ sở: Thi Lốp {b.name}"]
    if b.address:
        parts.append(f"Địa chỉ: {b.address}")
    if b.phone:
        parts.append(f"Điện thoại: {b.phone}")
    if b.hours:
        parts.append(f"Giờ mở cửa: {b.hours}")
    return "\n".join(parts)


@router.post("/quote")
def quote(body: QuoteIn, p: Principal = Depends(web_user), db: Session = Depends(get_db)):
    if not body.vehicle.strip():
        raise HTTPException(400, "Chưa nhập loại xe.")
    services = [s.strip() for s in body.services if s and s.strip()]
    b = db.get(Branch, body.branch_id) if body.branch_id else None
    if body.branch_id and b is None:
        raise HTTPException(400, "Không tìm thấy cơ sở.")
    query = " ".join(x for x in [body.vehicle, body.tire_size or "", *services, body.notes or ""] if x)
    results = _search(db, p, f"Giá và chính sách: {query}", k=5, public_only=True, feature="quote")
    lines = [f"Khách: {body.customer or 'không rõ tên'}", f"Xe: {body.vehicle}"]
    if body.tire_size:
        lines.append(f"Cỡ lốp: {body.tire_size}")
    lines.append("Dịch vụ cần báo giá: " + (", ".join(services) if services else "chưa ghi rõ"))
    if body.notes:
        lines.append(f"Ghi chú của nhân viên: {body.notes}")
    if b:
        lines.append("\n" + branch_info(b))
    ref = rag.format_context(results) if results else "(không có — mọi giá để [GIÁ])"
    user = "\n".join(lines) + f"\n\nThông tin tham khảo (giá, chính sách):\n{ref}\n\nSoạn tin báo giá."
    r = _call(db, p, "quote", [{"role": "system", "content": QUOTE_SYSTEM}, {"role": "user", "content": user}],
              task_hint="easy", branch_id=b.id if b else None)
    return {"text": r.content.strip(), "cost_vnd": float(r.cost_vnd)}


# ================= Biên bản họp (chạy nền) =================
AUDIO_EXT = {".mp3", ".m4a", ".wav"}


@router.post("/meeting")
def meeting(file: UploadFile = File(...), p: Principal = Depends(web_user), db: Session = Depends(get_db)):
    filename = file.filename or "ghi-am.mp3"
    if Path(filename).suffix.lower() not in AUDIO_EXT:
        raise HTTPException(400, "Chỉ nhận file ghi âm MP3, M4A hoặc WAV.")
    data = _read_upload(file, 25 * 1024 * 1024, "25 MB — hãy cắt ghi âm dài thành nhiều phần")
    rel = storage.save_bytes(data, filename, folder="audio")
    j = Job(kind="meeting", payload={"path": rel, "filename": filename}, created_by=p.user.id)
    db.add(j)
    db.commit()
    return {"job_id": j.id}


# ================= Phân tích doanh thu =================
VALUE_WORDS = ("doanh thu", "doanh_thu", "revenue", "thanh tien", "thanh_tien", "tien", "amount", "total", "tong")
GROUP_WORDS = ("co so", "co_so", "chi nhanh", "chi_nhanh", "branch", "dich vu", "dich_vu", "service", "hang muc", "loai")
DATE_WORDS = ("ngay", "date", "thoi gian", "thang", "month", "time")
MAX_REVENUE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 200_000

REVENUE_SYSTEM = (
    "Bạn là trợ lý phân tích kinh doanh của Thi Lốp. Dựa DUY NHẤT vào bản tóm tắt số liệu JSON được cung cấp "
    "(đã tính sẵn, không tự tính lại khác đi), viết nhận xét tiếng Việt dạng gạch đầu dòng (4–7 ý): "
    "nhóm đóng góp lớn nhất/nhỏ nhất, mức độ tập trung, xu hướng theo tháng nếu có, 1–2 gợi ý hành động cụ thể. "
    "Không bịa số liệu ngoài JSON. Số tiền viết dạng 1.250.000đ."
)


def _numeric_series(s):
    """Cột chữ nhưng thực chất là số (ví dụ '1.200.000đ') → chuyển thành số; không được thì trả None."""
    import pandas as pd

    if pd.api.types.is_bool_dtype(s):
        return None
    if pd.api.types.is_numeric_dtype(s):
        return s
    nonnull = s.dropna()
    if nonnull.empty:
        return None
    sample = nonnull.astype(str).head(500)
    if not sample.str.fullmatch(r"\s*-?[\d.,\s]+\s*(đ|d|vnd|vnđ|VND|VNĐ)?\s*").mean() >= 0.8:
        return None
    return s.map(lambda v: to_number(v) if v is not None and str(v).strip() not in ("", "nan") else None).astype("float64")


def _pick(names: list[str], words: tuple[str, ...]) -> str | None:
    for w in words:
        for n in names:
            if w in _plain(n):
                return n
    return None


def analyze_revenue(df) -> dict[str, Any]:
    """Tự nhận cột giá trị, cột nhóm, cột ngày rồi tính tổng, tỷ trọng, theo tháng (toàn bộ bằng Python)."""
    import pandas as pd

    df = df.dropna(how="all")
    df.columns = [str(c).strip() for c in df.columns]
    cols = list(df.columns)

    numeric: dict[str, Any] = {}
    for c in cols:
        ns = _numeric_series(df[c])
        if ns is not None and ns.notna().any():
            numeric[c] = ns

    # Cột ngày: tên gợi ý ngày/tháng hoặc kiểu datetime
    date_col = None
    for c in cols:
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            date_col = c
            break
    if date_col is None:
        for c in cols:
            if any(w in _plain(c) for w in DATE_WORDS) and not any(w in _plain(c) for w in VALUE_WORDS):
                parsed = pd.to_datetime(df[c].astype(str), errors="coerce", dayfirst=True, format="mixed")
                if parsed.notna().mean() >= 0.8:
                    date_col = c
                    break

    not_money = re.compile(r"\b(so luong|sl|qty|quantity|stt|id|ma|nam|year)\b")
    value_candidates = [c for c in numeric if c != date_col and not not_money.search(_plain(c))]
    value_col = _pick(value_candidates, VALUE_WORDS) or _pick(list(numeric), VALUE_WORDS)
    if value_col is None and value_candidates:
        value_col = max(value_candidates, key=lambda c: float(numeric[c].abs().sum()))
    if value_col is None:
        raise HTTPException(400, "Không tìm thấy cột số tiền (doanh thu / thành tiền) trong file.")

    text_cols = []
    for c in cols:
        if c in (value_col, date_col) or c in numeric:
            continue
        vals = df[c].dropna().astype(str).str.strip()
        nun = vals[vals != ""].nunique()
        if 2 <= nun <= 50:
            text_cols.append(c)
    group_col = _pick(text_cols, GROUP_WORDS) or (text_cols[0] if text_cols else None)

    values = numeric[value_col].fillna(0)
    total = float(values.sum())
    groups: list[dict] = []
    if group_col:
        keys = df[group_col].astype(str).str.strip().where(df[group_col].notna(), "(trống)")
        agg = values.groupby(keys).sum().sort_values(ascending=False)
        groups = [{"name": str(k), "total": round(float(v), 2),
                   "share": round(float(v) * 100 / total, 1) if total else 0.0} for k, v in agg.items()]
    by_period: list[dict] = []
    if date_col:
        d = df[date_col] if pd.api.types.is_datetime64_any_dtype(df[date_col]) else \
            pd.to_datetime(df[date_col].astype(str), errors="coerce", dayfirst=True, format="mixed")
        per = values.groupby(d.dt.strftime("%Y-%m")).sum().sort_index()
        by_period = [{"period": str(k), "total": round(float(v), 2)} for k, v in per.items() if str(k) != "nan"]
    return {"columns": cols, "group_by": group_col, "value_col": value_col, "date_col": date_col,
            "rows": int(len(df)), "total": round(total, 2), "groups": groups, "by_period": by_period}


def read_table(data: bytes, filename: str):
    import pandas as pd

    ext = Path(filename).suffix.lower()
    try:
        if ext == ".csv":
            tmp = io.BytesIO(data)
            for enc in ("utf-8-sig", "cp1258", "latin-1"):
                try:
                    tmp.seek(0)
                    return pd.read_csv(tmp, sep=None, engine="python", encoding=enc, nrows=MAX_ROWS + 1)
                except UnicodeDecodeError:
                    continue
            raise HTTPException(400, "Không đọc được bảng mã của file CSV (hãy lưu dạng UTF-8).")
        return pd.read_excel(io.BytesIO(data), nrows=MAX_ROWS + 1)
    except HTTPException:
        raise
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"Không đọc được file bảng: {e}") from e


@router.post("/revenue")
def revenue(file: UploadFile = File(...), p: Principal = Depends(web_user), db: Session = Depends(get_db)):
    filename = file.filename or "doanh-thu.csv"
    if Path(filename).suffix.lower() not in (".csv", ".xlsx"):
        raise HTTPException(400, "Chỉ nhận file CSV hoặc Excel (.xlsx).")
    data = _read_upload(file, MAX_REVENUE_BYTES, "5 MB")
    df = read_table(data, filename)
    if len(df) > MAX_ROWS:
        raise HTTPException(400, "File quá nhiều dòng (tối đa 200.000 dòng).")
    if df.empty:
        raise HTTPException(400, "File không có dòng dữ liệu nào.")
    summary = analyze_revenue(df)
    brief = {k: summary[k] for k in ("group_by", "value_col", "rows", "total")}
    brief["groups"] = summary["groups"][:20]
    brief["by_period"] = summary["by_period"][-24:]
    r = _call(db, p, "revenue", [
        {"role": "system", "content": REVENUE_SYSTEM},
        {"role": "user", "content": "Tóm tắt số liệu doanh thu (tỷ trọng tính theo %):\n" + json.dumps(brief, ensure_ascii=False)},
    ], task_hint="writing")
    summary.pop("date_col", None)
    summary.pop("rows", None)
    return {**summary, "commentary": r.content.strip(), "cost_vnd": float(r.cost_vnd)}

