"""Hỏi tài liệu: đọc chữ từ file, cắt đoạn, tạo vector (qua llm.embed) và tìm đoạn gần nhất bằng pgvector."""
from __future__ import annotations

import logging
import re
from pathlib import Path

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from ..models import DocChunk, Document
from . import llm
from .principal import Principal
from .storage import abs_path

log = logging.getLogger("hub.rag")

ALLOWED_EXT = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx"}
CHUNK_SIZE = 900
CHUNK_OVERLAP = 120


class ExtractError(Exception):
    pass


# ---------------- Đọc chữ ----------------
def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_csv(path: Path, **kw):
    """Đọc CSV, tự đoán dấu phân cách và bảng mã."""
    import pandas as pd

    last: Exception | None = None
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            return pd.read_csv(path, sep=None, engine="python", encoding=enc, **kw)
        except UnicodeDecodeError as e:
            last = e
    raise ExtractError(f"Không đọc được file CSV: {last}")


def _rows_to_lines(df) -> str:
    """Mỗi dòng bảng thành 'cột: giá trị; cột: giá trị' để tìm kiếm được."""
    import pandas as pd

    cols = [str(c) for c in df.columns]
    lines = []
    for row in df.itertuples(index=False, name=None):
        parts = [f"{c}: {v}" for c, v in zip(cols, row, strict=False) if not (v is None or (isinstance(v, float) and pd.isna(v))) and str(v).strip()]
        if parts:
            lines.append("; ".join(parts))
    return "\n".join(lines)


def extract_pages(path: Path, filename: str) -> list[tuple[int, str]]:
    """Trả về [(số trang, chữ)]. Word/CSV/TXT không có trang thì coi là trang 1 (Excel: mỗi sheet một trang)."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(str(path))
        except Exception as e:  # noqa: BLE001
            raise ExtractError(f"Không mở được PDF: {e}") from e
        return [(i, (page.extract_text() or "")) for i, page in enumerate(reader.pages, 1)]
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return [(1, "\n".join(parts))]
    if ext == ".csv":
        return [(1, _rows_to_lines(read_csv(path)))]
    if ext == ".xlsx":
        import pandas as pd

        sheets = pd.read_excel(path, sheet_name=None)
        return [(i, f"[{name}]\n" + _rows_to_lines(df)) for i, (name, df) in enumerate(sheets.items(), 1)]
    if ext in (".txt", ".md"):
        return [(1, _read_text_file(path))]
    raise ExtractError(f"Chưa hỗ trợ loại file {ext}.")


# ---------------- Cắt đoạn ----------------
def _clean(text: str) -> str:
    text = text.replace("\r", "").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def chunk_pages(pages: list[tuple[int, str]], size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[tuple[int, str]]:
    """Cắt mỗi trang thành đoạn ~size ký tự, chồng lấn ~overlap; ưu tiên cắt ở xuống dòng / dấu chấm / khoảng trắng."""
    out: list[tuple[int, str]] = []
    for page, raw in pages:
        text = _clean(raw)
        n, start = len(text), 0
        while start < n:
            end = min(start + size, n)
            if end < n:
                lo = start + size // 2
                cut = text.rfind("\n", lo, end)
                if cut < 0:
                    cut = text.rfind(". ", lo, end)
                    cut = cut + 1 if cut >= 0 else -1
                if cut < 0:
                    cut = text.rfind(" ", lo, end)
                if cut > start:
                    end = cut
            piece = text[start:end].strip()
            if piece:
                out.append((page, piece))
            if end >= n:
                break
            nxt = max(end - overlap, start + 1)
            sp = text.find(" ", nxt, end)  # bắt đầu đoạn sau ở đầu một từ
            start = sp + 1 if 0 <= sp < end else nxt
    return out


# ---------------- Lập chỉ mục ----------------
def index_document(db: Session, doc_id: int) -> int:
    """Đọc file, cắt đoạn, tạo vector, lưu DocChunk. Trả về số đoạn. Lỗi thì ghi vào tài liệu rồi ném lại."""
    doc = db.get(Document, doc_id)
    if doc is None:
        log.info("Tài liệu %s đã bị xoá, bỏ qua", doc_id)
        return 0
    doc.status, doc.error = "indexing", ""
    db.commit()
    try:
        pages = extract_pages(abs_path(doc.path), doc.filename)
        chunks = chunk_pages(pages)
        if not chunks:
            raise ExtractError("Không đọc được chữ nào trong file (có thể là PDF ảnh scan).")
        vectors = llm.embed(db, None, [t for _, t in chunks], app="apps", feature="index")
        db.execute(delete(DocChunk).where(DocChunk.document_id == doc_id))
        db.add_all(DocChunk(document_id=doc_id, page=pg, text=t, embedding=v) for (pg, t), v in zip(chunks, vectors, strict=True))
        doc.status, doc.pages, doc.chunks = "ready", len(pages), len(chunks)
        db.commit()
        return len(chunks)
    except Exception as e:
        db.rollback()
        doc = db.get(Document, doc_id)
        if doc is not None:
            doc.status = "error"
            doc.error = e.message if isinstance(e, llm.GatewayError) else str(e)[:1000]
            db.commit()
        raise


# ---------------- Tìm kiếm ----------------
def search(
    db: Session,
    query: str,
    k: int = 6,
    public_only: bool = False,
    *,
    p: Principal | None = None,
    app: str = "apps",
    feature: str = "search",
) -> list[tuple[DocChunk, Document, float]]:
    """Trả về [(đoạn, tài liệu, khoảng cách cosine)] gần câu hỏi nhất, chỉ trong tài liệu đã sẵn sàng."""
    if not query.strip():
        return []
    base = select(Document.id).where(Document.status == "ready")
    if public_only:
        base = base.where(Document.public_kb.is_(True))
    if db.execute(base.limit(1)).first() is None:
        return []  # chưa có tài liệu → khỏi tốn tiền tạo vector
    vec = llm.embed(db, p, [query], app=app, feature=feature)[0]
    dist = DocChunk.embedding.cosine_distance(vec).label("distance")
    q = select(DocChunk, Document, dist).join(Document, Document.id == DocChunk.document_id).where(Document.status == "ready")
    if public_only:
        q = q.where(Document.public_kb.is_(True))
    rows = db.execute(q.order_by(dist).limit(k)).all()
    return [(c, d, float(x)) for c, d, x in rows]


def format_context(results: list[tuple[DocChunk, Document, float]]) -> str:
    """Ghép các đoạn thành khối [1], [2]... để đưa vào câu lệnh."""
    return "\n\n".join(f"[{i}] ({d.title}, trang {c.page})\n{c.text}" for i, (c, d, _) in enumerate(results, 1))
