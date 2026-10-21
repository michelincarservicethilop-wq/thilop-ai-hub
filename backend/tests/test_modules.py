"""Kiểm thử các module: Tài liệu, Ứng dụng AI, Studio, Hộp thư, Webhook (chạy chế độ MOCK_LLM)."""
import hashlib
import hmac
import io
import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.config import get_settings
from app.db import session_factory
from app.models import CustomerConversation, Document, Job, UsageLog
from tests.conftest import login

PRICE_DOC = (
    "BẢNG GIÁ DỊCH VỤ THI LỐP 2026\n\n"
    "Cân bằng động 4 bánh: 200.000đ. Đảo lốp: 100.000đ. Bơm khí Nitơ: miễn phí khi thay lốp.\n"
    + "Keo chống đinh giúp lốp tự vá vết thủng nhỏ dưới 6 mm, hiệu quả trong 2 năm. " * 15
    + "\nGiờ mở cửa: 7:30 đến 18:00 tất cả các ngày trong tuần.\n"
)


def _run_jobs():
    from app.worker import process_one

    n = 0
    with session_factory()() as db:
        while process_one(db):
            n += 1
    return n


def _upload_doc(client, h, name="bang-gia.txt", data=PRICE_DOC, public=True):
    r = client.post("/api/documents", headers=h, files={"file": (name, data.encode(), "text/plain")},
                    data={"title": "Bảng giá 2026", "public_kb": "true" if public else "false"})
    assert r.status_code == 200, r.text
    return r.json()


# ---------------- Tài liệu & Hỏi tài liệu ----------------
def test_document_index_and_ask(client):
    h = login(client)
    d = _upload_doc(client, h)
    assert d["status"] == "queued"
    assert _run_jobs() == 1
    docs = client.get("/api/documents", headers=h).json()
    assert docs[0]["status"] == "ready" and docs[0]["chunks"] >= 2 and docs[0]["pages"] == 1, docs
    r = client.post("/api/apps/ask", json={"question": "Cân bằng động bao nhiêu tiền?"}, headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["answer"] and body["sources"]
    s = body["sources"][0]
    assert s["document_id"] == d["id"] and s["title"] == "Bảng giá 2026" and s["page"] == 1 and len(s["excerpt"]) <= 200
    # đổi tên, tắt dùng cho khách, xoá
    r = client.patch(f"/api/documents/{d['id']}", json={"title": "Giá mới", "public_kb": False}, headers=h)
    assert r.json()["title"] == "Giá mới" and r.json()["public_kb"] is False
    assert client.delete(f"/api/documents/{d['id']}", headers=h).status_code == 200
    assert client.get("/api/documents", headers=h).json() == []


def test_document_rules(client):
    h = login(client)
    r = client.post("/api/documents", headers=h, files={"file": ("virus.exe", b"x", "application/octet-stream")})
    assert r.status_code == 400
    client.post("/api/admin/users", json={"email": "nv@test.vn", "name": "NV", "password": "matkhau-nv-1", "role": "staff", "branch_id": 1}, headers=h)
    sh = login(client, "nv@test.vn", "matkhau-nv-1")
    r = client.post("/api/documents", headers=sh, files={"file": ("a.txt", b"abc", "text/plain")})
    assert r.status_code == 403
    # chưa có tài liệu → trả lời rõ ràng, không tốn tiền
    r = client.post("/api/apps/ask", json={"question": "Giá lốp?"}, headers=sh).json()
    assert r["sources"] == [] and r["cost_vnd"] == 0


def test_chunking_keeps_pages():
    from app.services.rag import chunk_pages

    chunks = chunk_pages([(1, "a " * 1000), (2, "Trang hai ngắn.")])
    assert {p for p, _ in chunks} == {1, 2}
    assert all(len(t) <= 900 for _, t in chunks)
    assert chunks[-1] == (2, "Trang hai ngắn.")


# ---------------- Hoá đơn ----------------
def test_invoice_parser_and_xlsx():
    from app.api.apps import invoice_xlsx, parse_invoice

    raw = """Đây là kết quả:
```json
{"fields": {"supplier": "Công ty Lốp Việt", "invoice_no": "0001234", "date": "20/10/2026", "total": "2.420.000", "tax": 220000},
 "items": [{"name": "Lốp 205/55R16", "quantity": 2, "unit_price": "1.000.000đ", "amount": null},
           {"name": "Van lốp", "quantity": "4", "unit_price": 50000, "amount": 200000}]}
```"""
    inv = parse_invoice(raw)
    assert inv["fields"]["supplier"] == "Công ty Lốp Việt"
    assert inv["fields"]["total"] == 2420000 and inv["fields"]["tax"] == 220000
    assert inv["items"][0]["amount"] == 2000000  # tự tính số lượng × đơn giá
    assert inv["items"][1]["quantity"] == 4
    assert parse_invoice("không có json")["items"] == []

    wb = load_workbook(io.BytesIO(invoice_xlsx(inv["fields"], inv["items"])))
    ws = wb.active
    header = next(r for r in ws.iter_rows() if r[0].value == "STT")
    assert header[1].value == "Tên hàng" and header[1].font.bold
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "Lốp 205/55R16" in values and 2200000 in values and 2420000 in values


def test_invoice_endpoint(client):
    h = login(client)
    r = client.post("/api/apps/invoice", headers=h, files={"file": ("hd.png", b"\x89PNG\r\n\x1a\nfake", "image/png")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["xlsx_url"].startswith("/api/files/exports/") and "fields" in body and "items" in body
    assert client.get(body["xlsx_url"], headers=h).status_code == 200
    r = client.post("/api/apps/invoice", headers=h, files={"file": ("hd.gif", b"GIF89a", "image/gif")})
    assert r.status_code == 400


# ---------------- Doanh thu ----------------
def test_revenue_csv(client):
    h = login(client)
    csv = (
        "Ngày,Cơ sở,Dịch vụ,Doanh thu\n"
        "05/09/2026,Lê Văn Lương,Thay lốp,1.000.000\n"
        "06/09/2026,Phạm Văn Đồng,Thay lốp,500.000\n"
        "02/10/2026,Lê Văn Lương,Ắc quy,2.000.000\n"
        "15/10/2026,Phạm Văn Đồng,Rửa xe,500.000\n"
    )
    r = client.post("/api/apps/revenue", headers=h, files={"file": ("dt.csv", csv.encode(), "text/csv")})
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["value_col"] == "Doanh thu" and b["group_by"] == "Cơ sở"
    assert b["total"] == 4000000
    g = {x["name"]: x for x in b["groups"]}
    assert g["Lê Văn Lương"]["total"] == 3000000 and g["Lê Văn Lương"]["share"] == 75.0
    assert g["Phạm Văn Đồng"]["share"] == 25.0
    assert b["by_period"] == [{"period": "2026-09", "total": 1500000}, {"period": "2026-10", "total": 2500000}]
    assert b["commentary"] and b["columns"] == ["Ngày", "Cơ sở", "Dịch vụ", "Doanh thu"]


# ---------------- Biên bản họp ----------------
def test_meeting_job(client):
    h = login(client)
    r = client.post("/api/apps/meeting", headers=h, files={"file": ("hop.mp3", b"ID3fake-audio", "audio/mpeg")})
    assert r.status_code == 200, r.text
    jid = r.json()["job_id"]
    assert client.get(f"/api/jobs/{jid}", headers=h).json()["status"] == "queued"
    _run_jobs()
    j = client.get(f"/api/jobs/{jid}", headers=h).json()
    assert j["status"] == "done", j
    assert set(j["result"]) == {"transcript", "summary", "decisions", "actions"} and j["result"]["transcript"]


# ---------------- Studio ----------------
def test_studio_script(client):
    h = login(client)
    r = client.post("/api/studio/script", json={"service": "Keo chống đinh", "platform": "tiktok", "length_sec": 30, "branch_id": 1}, headers=h)
    assert r.status_code == 200, r.text
    b = r.json()
    assert isinstance(b["script"], str) and b["script"]
    assert isinstance(b["caption"], str) and isinstance(b["hashtags"], list) and "cost_vnd" in b
    assert client.post("/api/studio/script", json={"service": "x", "platform": "youtube"}, headers=h).status_code == 400


def test_studio_hashtags_and_script_json():
    from app.api.studio import _hashtags, _script_text

    assert _hashtags(["thilop", "#Lop Xe", "#thilop"]) == ["#thilop", "#LopXe"]
    assert "Cảnh: cận lốp" in _script_text([{"time": "0-3s", "Cảnh": "cận lốp"}])


def test_tts_mock_and_cache(client):
    h = login(client)
    voices = client.get("/api/studio/voices", headers=h).json()
    ids = [v["id"] for v in voices]
    assert "fpt:banmai" in ids and "openai:nova" in ids
    assert all(v["available"] is False for v in voices)  # chưa cài key trong test
    assert "06/01/2027" in next(v for v in voices if v["id"] == "openai:nova")["label"]
    r = client.post("/api/studio/tts", json={"text": "Thi Lốp xin chào quý khách", "voice": "fpt:banmai"}, headers=h)
    assert r.status_code == 200, r.text
    first = r.json()
    assert first["url"].startswith("/api/files/tts/") and first["chars"] == 26 and first["cached"] is False
    assert client.get(first["url"] + "?token=" + h["Authorization"][7:]).status_code == 200
    second = client.post("/api/studio/tts", json={"text": "Thi Lốp xin chào quý khách", "voice": "fpt:banmai"}, headers=h).json()
    assert second["cached"] is True and second["cost_vnd"] == 0 and second["url"] == first["url"]
    assert client.post("/api/studio/tts", json={"text": "a" * 5001, "voice": "fpt:banmai"}, headers=h).status_code == 400
    assert client.post("/api/studio/tts", json={"text": "a", "voice": "fpt:khongco"}, headers=h).status_code == 400


def test_tts_quota(client):
    h = login(client)
    client.put("/api/admin/settings", json={"company_monthly_limit_vnd": 1000}, headers=h)
    with session_factory()() as db:  # giả lập đã tiêu hết hạn mức tháng
        db.add(UsageLog(app="chat", cost_vnd=1000))
        db.commit()
    r = client.post("/api/studio/tts", json={"text": "xin chào", "voice": "fpt:banmai"}, headers=h)
    assert r.status_code == 402


# ---------------- Báo giá ----------------
def test_quote(client):
    h = login(client)
    r = client.post("/api/apps/quote", json={"customer": "anh Nam", "vehicle": "Toyota Vios 2020", "tire_size": "185/60R15",
                                             "services": ["Thay 2 lốp", "Cân bằng động"], "branch_id": 1}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["text"] and r.json()["cost_vnd"] > 0


# ---------------- Hộp thư ----------------
def test_inbox_bot_replies(client):
    h = login(client)
    r = client.post("/api/inbox/simulate", json={"channel": "messenger", "external_user_id": "u1", "name": "Anh Minh",
                                                 "text": "Cơ sở Lê Văn Lương mở cửa mấy giờ?"}, headers=h)
    assert r.status_code == 200, r.text
    b = r.json()
    assert len(b["replies"]) == 1 and b["replies"][0]["sender"] == "bot"
    assert b["replies"][0]["text"].startswith("Chào anh/chị")  # lời chào lần đầu
    conv = client.get(f"/api/inbox/conversations/{b['conversation_id']}", headers=h).json()
    assert conv["status"] == "bot" and conv["branch_name"] == "Lê Văn Lương"
    assert conv["window"]["bot_can_reply"] is True and "giờ để bot trả lời" in conv["window"]["note"]
    assert [m["sender"] for m in conv["messages"]] == ["customer", "bot"]
    # tin thứ hai: không chào lại
    r2 = client.post("/api/inbox/simulate", json={"channel": "messenger", "external_user_id": "u1", "name": "Anh Minh",
                                                  "text": "Có đảo lốp không?"}, headers=h).json()
    assert not r2["replies"][0]["text"].startswith("Chào anh/chị")
    lst = client.get("/api/inbox/conversations", headers=h).json()
    assert lst[0]["last_message"] and lst[0]["channel"] == "messenger" and lst[0]["unread_human"] is False


def test_inbox_complaint_handoff(client):
    h = login(client)
    b = client.post("/api/inbox/simulate", json={"channel": "zalo", "external_user_id": "z1", "name": "Chị Lan",
                                                 "text": "làm ăn kiểu gì vậy, tôi muốn khiếu nại"}, headers=h).json()
    senders = [m["sender"] for m in b["replies"]]
    assert senders == ["bot", "system"]
    assert "chuyển cho nhân viên" in b["replies"][0]["text"]
    assert b["replies"][1]["text"].startswith("Chuyển cho nhân viên · ")
    conv = client.get(f"/api/inbox/conversations/{b['conversation_id']}", headers=h).json()
    assert conv["status"] == "human" and conv["handoff_reason"] == "complaint"
    # khách nhắn tiếp → bot im lặng
    again = client.post("/api/inbox/simulate", json={"channel": "zalo", "external_user_id": "z1", "name": "Chị Lan", "text": "Alo?"}, headers=h).json()
    assert again["replies"] == []
    lst = client.get("/api/inbox/conversations?status=human", headers=h).json()
    assert len(lst) == 1 and lst[0]["unread_human"] is True
    ov = client.get("/api/overview", headers=h).json()
    assert ov["inbox"]["handoffs"] == 1


def test_inbox_takeover_reply_release(client):
    h = login(client)
    cid = client.post("/api/inbox/simulate", json={"channel": "messenger", "external_user_id": "u9", "name": "Khách A",
                                                   "text": "Thay ắc quy hết bao lâu?"}, headers=h).json()["conversation_id"]
    c = client.post(f"/api/inbox/conversations/{cid}/takeover", headers=h).json()
    assert c["status"] == "human"
    m = client.post(f"/api/inbox/conversations/{cid}/reply", json={"text": "Dạ khoảng 15 phút ạ"}, headers=h).json()
    assert m["sender"] == "staff" and m["staff_name"] and m["sent"] is False  # chưa nối kênh → không gửi thật
    c = client.post(f"/api/inbox/conversations/{cid}/release", headers=h).json()
    assert c["status"] == "bot"
    c = client.post(f"/api/inbox/conversations/{cid}/branch", json={"branch_id": 2}, headers=h).json()
    assert c["branch_name"] == "Phạm Văn Đồng"
    c = client.post(f"/api/inbox/conversations/{cid}/done", headers=h).json()
    assert c["status"] == "done"
    # nhân viên trả lời khi bot đang trực → tự chuyển sang chế độ người
    client.post(f"/api/inbox/conversations/{cid}/release", headers=h)
    client.post(f"/api/inbox/conversations/{cid}/reply", json={"text": "Em chào anh"}, headers=h)
    assert client.get(f"/api/inbox/conversations/{cid}", headers=h).json()["status"] == "human"
    # nhân viên cơ sở 1 không thấy hội thoại của cơ sở 2
    client.post("/api/admin/users", json={"email": "nv1@test.vn", "name": "NV1", "password": "matkhau-nv-1", "role": "staff", "branch_id": 1}, headers=h)
    sh = login(client, "nv1@test.vn", "matkhau-nv-1")
    assert client.get(f"/api/inbox/conversations/{cid}", headers=sh).status_code == 404
    assert client.get("/api/inbox/conversations", headers=sh).json() == []
    assert client.post("/api/inbox/simulate", json={"text": "x"}, headers=sh).status_code == 403


def test_inbox_settings(client):
    h = login(client)
    s = client.get("/api/inbox/settings", headers=h).json()
    assert s["rules"]["complaint"] is True and s["bot_model"] == "auto"
    s2 = client.put("/api/inbox/settings", json={"rules": {"complaint": False}, "night_reply": False}, headers=h).json()
    assert s2["rules"] == {"complaint": False, "offprice": True, "askhuman": True, "unsure": True} and s2["night_reply"] is False
    assert client.put("/api/inbox/settings", json={"khong_co": 1}, headers=h).status_code == 400
    assert "zalo_tokens" not in client.get("/api/admin/settings", headers=h).json()["values"]
    # tắt luật phàn nàn → bot tự trả lời, không chuyển người
    client.put("/api/inbox/settings", json={"night_reply": True, "messenger_enabled": True}, headers=h)
    b = client.post("/api/inbox/simulate", json={"channel": "messenger", "external_user_id": "k", "name": "K", "text": "tôi muốn khiếu nại"}, headers=h).json()
    assert [m["sender"] for m in b["replies"]] == ["bot"]
    # tắt kênh → không trả lời
    client.put("/api/inbox/settings", json={"messenger_enabled": False}, headers=h)
    b = client.post("/api/inbox/simulate", json={"channel": "messenger", "external_user_id": "k2", "name": "K2", "text": "Giá lốp?"}, headers=h).json()
    assert b["replies"] == []


# ---------------- Webhook ----------------
@pytest.fixture()
def messenger_env(monkeypatch):
    monkeypatch.setenv("MESSENGER_VERIFY_TOKEN", "verify-123")
    monkeypatch.setenv("MESSENGER_APP_SECRET", "app-secret-xyz")
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def test_messenger_webhook(client, messenger_env):
    r = client.get("/webhooks/messenger", params={"hub.mode": "subscribe", "hub.verify_token": "verify-123", "hub.challenge": "987654"})
    assert r.status_code == 200 and r.text == "987654"
    assert client.get("/webhooks/messenger", params={"hub.mode": "subscribe", "hub.verify_token": "sai", "hub.challenge": "1"}).status_code == 403

    payload = {"object": "page", "entry": [{"id": "PAGE", "time": 1, "messaging": [
        {"sender": {"id": "PSID-1"}, "recipient": {"id": "PAGE"}, "timestamp": 1, "message": {"mid": "m-1", "text": "Giá đảo lốp bao nhiêu?"}},
        {"sender": {"id": "PAGE"}, "recipient": {"id": "PSID-1"}, "message": {"mid": "m-2", "text": "echo", "is_echo": True}},
        {"sender": {"id": "PSID-1"}, "recipient": {"id": "PAGE"}, "delivery": {"mids": ["m-0"]}},
    ]}]}
    raw = json.dumps(payload).encode()
    assert client.post("/webhooks/messenger", content=raw, headers={"Content-Type": "application/json",
                                                                    "X-Hub-Signature-256": "sha256=sai"}).status_code == 403
    sig = "sha256=" + hmac.new(b"app-secret-xyz", raw, hashlib.sha256).hexdigest()
    r = client.post("/webhooks/messenger", content=raw, headers={"Content-Type": "application/json", "X-Hub-Signature-256": sig})
    assert r.status_code == 200
    # Meta gửi lại cùng tin → không tạo trùng
    client.post("/webhooks/messenger", content=raw, headers={"Content-Type": "application/json", "X-Hub-Signature-256": sig})
    with session_factory()() as db:
        conv = db.execute(select(CustomerConversation).where(CustomerConversation.external_user_id == "PSID-1")).scalar_one()
        assert conv.channel == "messenger"
    h = login(client)
    msgs = client.get(f"/api/inbox/conversations/{conv.id}", headers=h).json()["messages"]
    assert [m["sender"] for m in msgs] == ["customer", "bot"]
    assert msgs[1]["sent"] is False  # chế độ thử: không gọi Graph API


def test_zalo_webhook(client):
    payload = {"app_id": "1", "event_name": "user_send_text", "timestamp": "1761100000000",
               "sender": {"id": "ZUID-7"}, "recipient": {"id": "OA"}, "message": {"msg_id": "zm-1", "text": "Cơ sở Phạm Văn Đồng ở đâu?"}}
    r = client.post("/webhooks/zalo", json=payload)
    assert r.status_code == 200
    client.post("/webhooks/zalo", json={"event_name": "user_seen_message", "sender": {"id": "ZUID-8"}})
    with session_factory()() as db:
        convs = db.execute(select(CustomerConversation)).scalars().all()
        assert [(c.channel, c.external_user_id) for c in convs] == [("zalo", "ZUID-7")]
        assert convs[0].branch.name == "Phạm Văn Đồng"


def test_zalo_signature(client, monkeypatch):
    monkeypatch.setenv("ZALO_APP_ID", "app1")
    monkeypatch.setenv("ZALO_OA_SECRET_KEY", "oa-secret")
    get_settings.cache_clear()
    try:
        payload = {"event_name": "user_send_text", "timestamp": "123", "sender": {"id": "Z9"}, "message": {"msg_id": "a", "text": "hi"}}
        raw = json.dumps(payload).encode()
        assert client.post("/webhooks/zalo", content=raw, headers={"X-ZEvent-Signature": "mac=sai"}).status_code == 403
        mac = hashlib.sha256(("app1" + raw.decode() + "123" + "oa-secret").encode()).hexdigest()
        assert client.post("/webhooks/zalo", content=raw, headers={"X-ZEvent-Signature": "mac=" + mac}).status_code == 200
    finally:
        monkeypatch.undo()
        get_settings.cache_clear()


def test_worker_marks_bad_job_error(client):
    with session_factory()() as db:
        db.add(Job(kind="khong_ro", payload={}))
        db.add(Document(title="x", filename="x.txt", path="documents/khong-co.txt", status="queued"))
        db.commit()
        doc_id = db.execute(select(Document.id)).scalar_one()
        db.add(Job(kind="index_document", payload={"document_id": doc_id}))
        db.commit()
    assert _run_jobs() == 2
    with session_factory()() as db:
        jobs = db.execute(select(Job).order_by(Job.id)).scalars().all()
        assert [j.status for j in jobs] == ["error", "error"]
        assert db.get(Document, doc_id).status == "error"
